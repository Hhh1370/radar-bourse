#!/usr/bin/env python3
"""Extract public Bourse-Trader/Telegram Codal references into a reviewable XLSX.

This is a discovery workbook, not an audited financial parser. It preserves original
URLs, post text, dates, and keyword evidence. Unreachable Codal pages are never
interpreted as zero activity.
"""
from __future__ import annotations

import argparse
import html
import io
import json
import re
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    raise SystemExit("Missing dependency openpyxl; install it before running this script.") from exc

ROOT = Path(__file__).resolve().parents[1]
OUT_XLSX = ROOT / "data" / "codal_disclosures_latest.xlsx"
OUT_JSON = ROOT / "data" / "codal_disclosures_latest.json"

BOURSE_URL = "https://bourse-trader.ir/index.php?task=codal-reports"
TELEGRAM_BASE = "https://t.me/s/codalboursetrader"
USER_AGENT = (
    "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/130.0.0.0 Safari/537.36"
)
PAGE_TIMEOUT = 18
MAX_PAGE_BYTES = 8 * 1024 * 1024
ATTACHMENT_TIMEOUT = 15
MAX_ATTACHMENT_BYTES = 8 * 1024 * 1024
MAX_ATTACHMENT_TESTS = 3
MAX_ATTACHMENT_ROWS_PER_SHEET = 2500
DEFAULT_TELEGRAM_PAGES = 8

PROJECT_KEYWORDS = (
    "طرح توسعه", "طرح احداث", "پروژه", "پیشرفت فیزیکی", "پیشرفت پروژه",
    "افزایش ظرفیت", "ظرفیت تولید", "بهره برداری", "بهره‌برداری",
    "راه اندازی", "راه‌اندازی", "تولید آزمایشی", "تکمیل طرح", "توسعه خطوط",
    "عملیات اجرایی", "سرمایه گذاری در طرح", "سرمایه‌گذاری در طرح",
    "فاز توسعه", "ماشین آلات طرح", "ماشین‌آلات طرح", "ظرفیت اسمی",
    "گزارش توجیهی طرح", "پیشرفت طرح", "طرح سرمایه‌گذاری",
)
FINANCIAL_KEYWORDS = (
    "صورت های مالی", "صورت‌های مالی", "صورت مالی", "گزارش تفسیری مدیریت",
    "گزارش حسابرسی", "اظهارنظر حسابرس", "ترازنامه", "سود و زیان",
    "جریان وجوه نقد", "حساب‌های مالی",
)
MONTHLY_KEYWORDS = (
    "گزارش فعالیت ماهانه", "فعالیت ماهانه", "تولید و فروش", "فروش ماهانه",
    "درآمد ماهانه", "عملکرد ماهانه",
)
IMPORTANT_KEYWORDS = (
    "افشای اطلاعات بااهمیت", "اطلاعات بااهمیت", "ماده 13", "ماده ۱۳",
    "ماده 129", "ماده ۱۲۹", "دریافت مجوز", "افزایش سرمایه",
    "انعقاد قرارداد", "قرارداد مهم", "مزایده", "مناقصه",
)
CAPITAL_KEYWORDS = ("افزایش سرمایه", "حق تقدم", "سلب حق تقدم", "افزایش سرمایه از")
CONTRACT_KEYWORDS = ("انعقاد قرارداد", "قرارداد مهم", "مناقصه", "مزایده")

def now_utc() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def clean_text(value: object) -> str:
    text = html.unescape(str(value or ""))
    text = text.replace("\u200c", " ").replace("\u200f", "").replace("\u200e", "")
    text = text.replace("\xa0", " ")
    return re.sub(r"\s+", " ", text).strip()


def normalize(value: object) -> str:
    text = clean_text(value)
    table = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")
    text = text.translate(table).replace("ي", "ی").replace("ك", "ک")
    return re.sub(r"[\s‌\-_]+", "", text).casefold()


def canonical_url(url: str) -> str:
    url = html.unescape(clean_text(url))
    parsed = urllib.parse.urlsplit(url)
    if parsed.scheme not in ("http", "https") or not parsed.netloc:
        return url
    return urllib.parse.urlunsplit((parsed.scheme.lower(), parsed.netloc.lower(), parsed.path, parsed.query, ""))


def url_host(url: str) -> str:
    return (urllib.parse.urlsplit(url).hostname or "").lower()


def is_codal_host(url: str) -> bool:
    host = url_host(url)
    return host == "codal.ir" or host.endswith(".codal.ir")


def is_codal_reference(url: str) -> bool:
    if not is_codal_host(url):
        return False
    parsed = urllib.parse.urlsplit(url)
    path = parsed.path.lower()
    query = parsed.query.lower()
    return (
        "reportlist.aspx" in path or "decision.aspx" in path or
        "/reports/" in path or "letterserial=" in query or
        "symbol=" in query
    )


def is_report_list(url: str) -> bool:
    return "reportlist.aspx" in urllib.parse.urlsplit(url).path.lower()


def is_attachment_candidate(url: str) -> bool:
    parsed = urllib.parse.urlsplit(url)
    path = parsed.path.lower()
    host = (parsed.hostname or "").lower()
    return (
        host == "excel.codal.ir" or host.endswith(".excel.codal.ir") or
        path.endswith(".xlsx") or path.endswith(".xls")
    )


def request_url(url: str, timeout: int = PAGE_TIMEOUT, max_bytes: int = MAX_PAGE_BYTES) -> dict:
    request = urllib.request.Request(
        url,
        headers={
            "User-Agent": USER_AGENT,
            "Accept": "text/html,application/xhtml+xml,application/json,*/*;q=0.8",
            "Accept-Language": "fa-IR,fa;q=0.9,en;q=0.7",
            "Referer": "https://bourse-trader.ir/",
        },
    )
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            body = response.read(max_bytes + 1)
            too_large = len(body) > max_bytes
            if too_large:
                body = body[:max_bytes]
            return {
                "url": url, "http_status": response.status,
                "content_type": response.headers.get("Content-Type", ""),
                "body": body, "body_bytes": len(body), "too_large": too_large,
                "error": "",
            }
    except urllib.error.HTTPError as exc:
        try:
            excerpt = exc.read(500).decode("utf-8", "replace")
        except Exception:
            excerpt = ""
        return {
            "url": url, "http_status": exc.code,
            "content_type": exc.headers.get("Content-Type", "") if exc.headers else "",
            "body": b"", "body_bytes": 0, "too_large": False,
            "error": clean_text(excerpt)[:240] or str(exc),
        }
    except Exception as exc:
        return {
            "url": url, "http_status": "ERROR",
            "content_type": "", "body": b"", "body_bytes": 0,
            "too_large": False, "error": f"{type(exc).__name__}: {exc}"[:300],
        }


class TelegramParser(HTMLParser):
    """Parse Telegram's public preview by actual message-wrap containers."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.posts = []
        self.current = None
        self.div_depth = 0
        self.anchor_stack = []

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        classes = attrs.get("class", "")
        if self.current is None and tag.lower() == "div" and "tgme_widget_message_wrap" in classes:
            self.current = {
                "post_id": "", "date": "", "text_parts": [], "links": [],
            }
            self.div_depth = 1
            return
        if self.current is None:
            return
        if tag.lower() == "div":
            self.div_depth += 1
        data_post = attrs.get("data-post", "")
        if data_post and "/" in data_post:
            self.current["post_id"] = data_post.rsplit("/", 1)[-1]
        if tag.lower() == "time":
            self.current["date"] = attrs.get("datetime", "") or attrs.get("title", "")
        if tag.lower() == "a" and attrs.get("href"):
            self.anchor_stack.append({"href": html.unescape(attrs["href"]), "parts": []})

    def handle_data(self, data):
        if self.current is None:
            return
        if clean_text(data):
            self.current["text_parts"].append(clean_text(data))
            if self.anchor_stack:
                self.anchor_stack[-1]["parts"].append(clean_text(data))

    def handle_endtag(self, tag):
        if self.current is None:
            return
        if tag.lower() == "a" and self.anchor_stack:
            anchor = self.anchor_stack.pop()
            self.current["links"].append({
                "url": anchor["href"],
                "title": clean_text(" ".join(anchor["parts"])),
            })
        if tag.lower() == "div":
            self.div_depth -= 1
            if self.div_depth <= 0:
                post = {
                    "post_id": clean_text(self.current["post_id"]),
                    "date": clean_text(self.current["date"]),
                    "text": clean_text(" ".join(self.current["text_parts"])),
                    "links": self.current["links"],
                }
                if post["text"] or post["links"]:
                    self.posts.append(post)
                self.current = None
                self.div_depth = 0
                self.anchor_stack = []


class PageAnchorParser(HTMLParser):
    """Collect links and their nearest table/list/article context from HTML."""

    VOID_TAGS = {"area", "base", "br", "col", "embed", "hr", "img", "input",
                 "link", "meta", "param", "source", "track", "wbr"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.stack = []
        self.links = []

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag.lower() in self.VOID_TAGS:
            return
        self.stack.append({"tag": tag.lower(), "attrs": attrs, "parts": []})

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)
        self.handle_endtag(tag)

    def handle_data(self, data):
        if not clean_text(data):
            return
        for node in self.stack:
            node["parts"].append(clean_text(data))

    def handle_endtag(self, tag):
        tag = tag.lower()
        idx = next((i for i in range(len(self.stack) - 1, -1, -1)
                    if self.stack[i]["tag"] == tag), None)
        if idx is None:
            return
        node = self.stack[idx]
        if tag == "a" and node["attrs"].get("href"):
            ancestors = self.stack[:idx]
            context = ""
            for ancestor in reversed(ancestors):
                cls = ancestor["attrs"].get("class", "")
                if (
                    ancestor["tag"] in {"tr", "li", "article"} or
                    re.search(r"(?:item|post|news|card|row|entry)", cls, re.I)
                ):
                    context = clean_text(" ".join(ancestor["parts"]))
                    if context:
                        break
            if not context and ancestors:
                context = clean_text(" ".join(ancestors[-1]["parts"]))
            self.links.append({
                "url": html.unescape(node["attrs"]["href"]),
                "title": clean_text(" ".join(node["parts"])),
                "context": context[:1800],
            })
        del self.stack[idx:]


def filter_navigation_context(context: str, link_title: str = "") -> str:
    """Avoid assigning sidebar/navigation keywords to the linked disclosure."""
    text = clean_text(context)
    nav_phrases = (
        "سابقه معاملات", "ریز معاملات امروز", "سهامداران عمده",
        "افزایش سرمایه", "تعدیل قیمت", "تابلو عرضه و تقاضا",
        "اندیکاتورها", "در TSETMC", "در کدال", "رتبه بندی بنیادی",
    )
    normalized = normalize(text)
    nav_hits = sum(1 for phrase in nav_phrases if normalize(phrase) in normalized)
    if nav_hits >= 3:
        return clean_text(link_title)
    return text or clean_text(link_title)


def extract_symbol(url: str, text: str) -> str:
    query = urllib.parse.parse_qs(urllib.parse.urlsplit(url).query)
    for key in ("Symbol", "symbol", "symbolName", "symbolname"):
        for value in query.get(key, []):
            value = clean_text(value)
            if value and len(value) <= 30:
                return value
    match = re.search(r"(?:نماد|شرکت)\s*[:：=\-]\s*([آ-یA-Za-z0-9_‌\-]{1,30})", text)
    if match:
        return clean_text(match.group(1))
    ignored = {
        "کدال", "بورس", "فرابورس", "بورس_تریدر", "صورت_مالی",
        "گزارش", "افشا", "سهام", "بازار", "طرح_توسعه",
    }
    for tag in re.findall(r"#([آ-یA-Za-z0-9_‌\-]{2,30})", text):
        normalized_tag = normalize(tag)
        if normalized_tag not in {normalize(item) for item in ignored} and not any(
            normalize(word) in normalized_tag for word in ("افشا", "گزارش", "صورت", "بورس", "کدال")
        ):
            return clean_text(tag).replace("_", " ")
    return ""


def find_terms(text: str, terms: tuple[str, ...]) -> list[str]:
    haystack = normalize(text)
    found = []
    for term in terms:
        if normalize(term) in haystack and term not in found:
            found.append(term)
    return found


def classify_text(text: str, link: str = "") -> str:
    if is_report_list(link):
        return "فهرست گزارش‌های نماد"
    if find_terms(text, FINANCIAL_KEYWORDS):
        return "صورت‌های مالی/حسابرسی"
    if find_terms(text, MONTHLY_KEYWORDS):
        return "گزارش فعالیت ماهانه/تولید و فروش"
    if find_terms(text, PROJECT_KEYWORDS):
        return "طرح توسعه/پروژه"
    if find_terms(text, CAPITAL_KEYWORDS):
        return "افزایش سرمایه"
    if find_terms(text, CONTRACT_KEYWORDS):
        return "قرارداد/مناقصه/مزایده"
    if find_terms(text, IMPORTANT_KEYWORDS):
        return "افشای اطلاعات بااهمیت"
    return "نیازمند دسته‌بندی دستی"


def raw_amount_mentions(text: str) -> str:
    pattern = re.compile(
        r"(?<!\w)([\d۰-۹٠-٩][\d۰-۹٠-٩,٬٫.]*\s*(?:(?:هزار|میلیون|میلیارد|تریلیون)\s*)?"
        r"(?:ریال|تومان|درصد|٪|%|تن|کیلوگرم|مترمربع|مگاوات|میلیون\s*دلار))",
        re.I,
    )
    values = []
    for match in pattern.finditer(text):
        value = clean_text(match.group(1))
        if value not in values:
            values.append(value)
    return "؛ ".join(values[:20])


def raw_dates(text: str) -> str:
    pattern = r"(?<!\d)(?:13|14)?[۰-۹0-9]{2,4}[/-][۰-۹0-9]{1,2}[/-][۰-۹0-9]{1,2}(?!\d)"
    values = []
    table = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")
    for value in re.findall(pattern, text):
        value = value.translate(table)
        if value not in values:
            values.append(value)
    return "؛ ".join(values[:10])


def telegram_link_for(post_id: str) -> str:
    if post_id.isdigit():
        return f"https://t.me/codalboursetrader/{post_id}"
    return "https://t.me/s/codalboursetrader"


def make_record(url: str, link_title: str, context: str, source: str,
                source_page: str, post_date: str = "", post_id: str = "",
                all_page_links: list[str] | None = None) -> dict:
    text = clean_text(" ".join(x for x in (link_title, context) if x))
    project_terms = find_terms(text, PROJECT_KEYWORDS)
    financial_terms = find_terms(text, FINANCIAL_KEYWORDS + MONTHLY_KEYWORDS)
    return {
        "url": canonical_url(url),
        "link_type": "فهرست نماد" if is_report_list(url) else "گزارش کدال",
        "link_title": clean_text(link_title),
        "symbol": extract_symbol(url, text),
        "date": clean_text(post_date) or raw_dates(text),
        "primary_category": classify_text(text, url),
        "project_keywords": "؛ ".join(project_terms),
        "financial_keywords": "؛ ".join(financial_terms),
        "raw_amount_mentions": raw_amount_mentions(text),
        "source_context": clean_text(context)[:1800],
        "source": source,
        "source_page": source_page,
        "post_id": post_id,
    }


def merge_records(records: list[dict]) -> list[dict]:
    merged = {}
    for row in records:
        key = row["url"] or f'{row["source"]}:{row["source_context"][:120]}'
        if key not in merged:
            merged[key] = dict(row)
            merged[key]["sources"] = [row["source"]]
            merged[key]["source_pages"] = [row["source_page"]]
            continue
        existing = merged[key]
        # Capture existing provenance before merging the incoming source.
        existing_has_telegram = "تلگرام بورس‌تریدر" in existing["sources"]
        row_is_telegram = row["source"] == "تلگرام بورس‌تریدر"
        if row["source"] not in existing["sources"]:
            existing["sources"].append(row["source"])
        if row["source_page"] and row["source_page"] not in existing["source_pages"]:
            existing["source_pages"].append(row["source_page"])
        # A Telegram post normally contains the actual disclosure description;
        # do not let a larger navigation/sidebar fragment from a secondary page
        # replace a shorter, relevant post context.
        if row_is_telegram and not existing_has_telegram:
            existing["source_context"] = row["source_context"]
        elif row_is_telegram == existing_has_telegram and len(row.get("source_context", "")) > len(existing.get("source_context", "")):
            existing["source_context"] = row["source_context"]
        if not existing.get("link_title") and row.get("link_title"):
            existing["link_title"] = row["link_title"]
        if not existing.get("symbol") and row.get("symbol"):
            existing["symbol"] = row["symbol"]
        if not existing.get("date") and row.get("date"):
            existing["date"] = row["date"]
        if not existing.get("project_keywords") and row.get("project_keywords"):
            existing["project_keywords"] = row["project_keywords"]
        if not existing.get("financial_keywords") and row.get("financial_keywords"):
            existing["financial_keywords"] = row["financial_keywords"]
        if existing.get("primary_category") == "نیازمند دسته‌بندی دستی" and row.get("primary_category") != "نیازمند دسته‌بندی دستی":
            existing["primary_category"] = row["primary_category"]
    for row in merged.values():
        row["source"] = "؛ ".join(row.pop("sources", []))
        row["source_page"] = "؛ ".join(row.pop("source_pages", []))
        # Recompute text-derived fields from the preferred, de-duplicated context.
        evidence = clean_text(" ".join((row.get("link_title", ""), row.get("source_context", ""))))
        row["primary_category"] = classify_text(evidence, row.get("url", ""))
        row["project_keywords"] = "؛ ".join(find_terms(evidence, PROJECT_KEYWORDS))
        row["financial_keywords"] = "؛ ".join(find_terms(evidence, FINANCIAL_KEYWORDS + MONTHLY_KEYWORDS))
        row["raw_amount_mentions"] = raw_amount_mentions(evidence)
        if not row.get("symbol"):
            row["symbol"] = extract_symbol(row.get("url", ""), evidence)
    return sorted(merged.values(), key=lambda item: (item.get("date", ""), item.get("symbol", ""), item.get("url", "")), reverse=True)


def fetch_telegram_pages(page_limit: int, source_status: list[dict]) -> list[dict]:
    posts_by_id = {}
    seen_urls = set()
    before = None
    previous_oldest = None

    for page_index in range(page_limit):
        page_url = TELEGRAM_BASE if before is None else f"{TELEGRAM_BASE}?before={before}"
        if page_url in seen_urls:
            break
        seen_urls.add(page_url)
        result = request_url(page_url)
        status_row = {
            "source_name": f"تلگرام بورس‌تریدر - صفحه {page_index + 1}",
            "url": page_url, "http_status": result["http_status"],
            "content_type": result["content_type"], "body_bytes": result["body_bytes"],
            "items_found": 0, "note": result["error"],
        }
        source_status.append(status_row)
        if result["http_status"] != 200 or not result["body"]:
            break
        parser = TelegramParser()
        try:
            parser.feed(result["body"].decode("utf-8", "replace"))
        except Exception as exc:
            status_row["note"] = f"HTML parse issue: {type(exc).__name__}: {exc}"
            break
        status_row["items_found"] = len(parser.posts)
        if not parser.posts:
            status_row["note"] = (status_row["note"] + "; " if status_row["note"] else "") + "HTTP 200 but no Telegram message blocks parsed"
            break

        numeric_ids = []
        for post in parser.posts:
            post_id = post.get("post_id", "")
            if post_id.isdigit():
                numeric_ids.append(int(post_id))
                key = post_id
            else:
                key = f'no-id:{post.get("date","")}:{post.get("text","")[:140]}'
            post["source_page"] = page_url
            post["permalink"] = telegram_link_for(post_id)
            posts_by_id[key] = post
        if not numeric_ids:
            status_row["note"] = "Messages parsed, but post ids were not detected; archive pagination stopped."
            break
        oldest = min(numeric_ids)
        if previous_oldest is not None and oldest >= previous_oldest:
            status_row["note"] = "Archive did not move to older post ids; stopped to prevent looping."
            break
        previous_oldest = oldest
        before = oldest - 1
        time.sleep(0.15)

    return list(posts_by_id.values())


def request_attachment_with_curl(url: str) -> dict:
    """Download a bounded Codal attachment with curl, which may handle this host better than urllib."""
    with tempfile.NamedTemporaryFile(prefix="codal-xlsx-", suffix=".bin") as target:
        command = [
            "curl", "--location", "--http1.1",
            "--connect-timeout", "7", "--max-time", str(ATTACHMENT_TIMEOUT),
            "--max-filesize", str(MAX_ATTACHMENT_BYTES), "--retry", "0",
            "--user-agent", USER_AGENT, "--header", "Accept: */*",
            "--silent", "--show-error", "--output", target.name,
            "--write-out", "%{http_code}\\t%{content_type}\\t%{url_effective}", url,
        ]
        try:
            proc = subprocess.run(
                command, capture_output=True, text=True,
                timeout=ATTACHMENT_TIMEOUT + 3, check=False,
            )
            stdout = (proc.stdout or "").strip()
            fields = stdout.split("\\t", 2)
            status = fields[0] if fields and fields[0] else "ERROR"
            try:
                status = int(status)
            except (TypeError, ValueError):
                pass
            content_type = fields[1] if len(fields) > 1 else ""
            effective_url = fields[2] if len(fields) > 2 else url
            target.flush()
            actual_size = Path(target.name).stat().st_size
            too_large = actual_size > MAX_ATTACHMENT_BYTES or proc.returncode == 63
            with open(target.name, "rb") as body_file:
                body = body_file.read(MAX_ATTACHMENT_BYTES + 1)
            err = clean_text(proc.stderr or "")
            if proc.returncode != 0 and not err:
                err = f"curl exited with code {proc.returncode}"
            if too_large:
                err = f"Attachment exceeded the {MAX_ATTACHMENT_BYTES}-byte safety limit."
            return {
                "url": effective_url or url,
                "http_status": status,
                "content_type": content_type,
                "body": body,
                "body_bytes": len(body),
                "too_large": too_large,
                "error": err[:300],
            }
        except subprocess.TimeoutExpired:
            target.flush()
            actual_size = Path(target.name).stat().st_size
            with open(target.name, "rb") as body_file:
                body = body_file.read(MAX_ATTACHMENT_BYTES + 1)
            return {
                "url": url, "http_status": "ERROR", "content_type": "",
                "body": body, "body_bytes": len(body),
                "too_large": actual_size > MAX_ATTACHMENT_BYTES,
                "error": f"curl timed out after {ATTACHMENT_TIMEOUT} seconds",
            }


def download_xlsx_attachments(candidates: list[dict], source_status: list[dict]):
    attachment_results = []
    extracted_rows = []
    unique = {}
    for candidate in candidates:
        url = canonical_url(candidate["url"])
        if url and is_attachment_candidate(url) and url not in unique:
            unique[url] = candidate
    selected = list(unique.items())[:MAX_ATTACHMENT_TESTS]

    for url, candidate in selected:
        result = request_attachment_with_curl(url)
        body = result["body"]
        is_zip = bool(body) and zipfile.is_zipfile(io.BytesIO(body))
        is_xlsx = False
        parse_note = ""
        sheet_count = 0
        extracted_count = 0

        if result["too_large"]:
            parse_note = f"فایل از سقف {MAX_ATTACHMENT_BYTES} بایت بزرگ‌تر است؛ فقط بخشی از پاسخ خوانده شد."
        elif result["http_status"] != 200 or not body:
            parse_note = result["error"] or "دانلود ناموفق/بدون بدنه"
        elif not is_zip:
            parse_note = "پاسخ دریافتی ZIP/XLSX معتبر نبود؛ احتمالاً صفحه خطا یا فرمت دیگر است."
        else:
            try:
                wb = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
                is_xlsx = True
                sheet_count = len(wb.sheetnames)
                for ws in wb.worksheets:
                    for row_no, cells in enumerate(ws.iter_rows(values_only=True), start=1):
                        vals = [clean_text(v) for v in cells]
                        if any(vals):
                            extracted_rows.append({
                                "attachment_url": url,
                                "sheet_name": ws.title,
                                "row_number": row_no,
                                "row_text": " | ".join(v for v in vals if v)[:4000],
                            })
                            extracted_count += 1
                        if extracted_count >= MAX_ATTACHMENT_ROWS_PER_SHEET:
                            parse_note = f"خواندن ردیف‌ها برای کنترل حجم در {MAX_ATTACHMENT_ROWS_PER_SHEET} ردیف متوقف شد."
                            break
                    if extracted_count >= MAX_ATTACHMENT_ROWS_PER_SHEET:
                        break
                wb.close()
                if not parse_note:
                    parse_note = "فایل XLSX خوانده شد؛ ارقام خام هستند و هنوز از نظر واحد/دوره مالی اعتبارسنجی نشده‌اند."
            except Exception as exc:
                parse_note = f"پاسخ ZIP بود اما بازکردن XLSX شکست خورد: {type(exc).__name__}: {exc}"[:400]

        attachment_results.append({
            "url": url, "host": url_host(url),
            "http_status": result["http_status"],
            "content_type": result["content_type"],
            "body_bytes": result["body_bytes"],
            "is_xlsx": "بله" if is_xlsx else "خیر/تأییدنشده",
            "sheet_count": sheet_count,
            "extracted_rows": extracted_count,
            "parse_note": parse_note,
            "source": candidate.get("source", ""),
            "source_context": candidate.get("context", "")[:400],
        })
        source_status.append({
            "source_name": "پیوست اکسل",
            "url": url, "http_status": result["http_status"],
            "content_type": result["content_type"], "body_bytes": result["body_bytes"],
            "items_found": extracted_count, "note": parse_note,
        })
        time.sleep(0.15)

    if len(unique) > len(selected):
        attachment_results.append({
            "url": "", "host": "", "http_status": "SKIPPED",
            "content_type": "", "body_bytes": 0, "is_xlsx": "خیر/تأییدنشده",
            "sheet_count": 0, "extracted_rows": 0,
            "parse_note": f"{len(unique) - len(selected)} پیوست دیگر برای محدودکردن زمان اجرا آزمایش نشد.",
            "source": "", "source_context": "",
        })
    return attachment_results, extracted_rows


def excel_safe(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False)
    text = str(value)
    # Avoid Excel interpreting scraped text as a formula.
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def write_sheet(wb, name: str, headers: list[str], rows: list[dict], key_map: list[str]):
    ws = wb.create_sheet(title=name)
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = "A2"
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30

    for row in rows:
        ws.append([excel_safe(row.get(key, "")) for key in key_map])
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    for col_index, header in enumerate(headers, start=1):
        sample_lengths = [len(str(ws.cell(r, col_index).value or "")) for r in range(1, min(ws.max_row, 80) + 1)]
        width = max([len(str(header)) + 2] + sample_lengths)
        if header in {"متن کامل پست", "متن زمینه", "ردیف استخراج‌شده", "یادداشت"}:
            width = min(max(width, 32), 70)
        else:
            width = min(max(width, 12), 48)
        ws.column_dimensions[get_column_letter(col_index)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    return ws


def build_workbook(payload: dict):
    wb = Workbook()
    ws = wb.active
    ws.title = "خلاصه"
    ws.sheet_view.rightToLeft = True
    summary = payload["summary"]
    ws.append(["شاخص", "مقدار", "توضیح"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for item in summary:
        ws.append([excel_safe(item.get("metric")), excel_safe(item.get("value")), excel_safe(item.get("note"))])
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 95
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    core_headers = [
        "نوع پیوند", "عنوان لینک", "نماد استخراج‌شده", "تاریخ انتشار/متن",
        "دسته‌بندی اولیه", "کلیدواژه‌های طرح توسعه", "کلیدواژه‌های مالی/ماهانه",
        "ارقام و واحدهای خام", "متن زمینه", "لینک اطلاعیه", "منبع استخراج", "صفحه/پست منبع",
    ]
    core_keys = [
        "link_type", "link_title", "symbol", "date", "primary_category",
        "project_keywords", "financial_keywords", "raw_amount_mentions",
        "source_context", "url", "source", "source_page",
    ]
    announcements = payload["announcements"]
    write_sheet(wb, "اطلاعیه‌ها", core_headers, announcements, core_keys)
    projects = [row for row in announcements if row.get("project_keywords")]
    write_sheet(wb, "طرح‌های توسعه", core_headers, projects, core_keys)
    finances = [row for row in announcements if row.get("financial_keywords") or "صورت‌های مالی" in row.get("primary_category", "") or "ماهانه" in row.get("primary_category", "")]
    write_sheet(wb, "صورت‌های مالی", core_headers, finances, core_keys)

    post_rows = []
    for post in payload["telegram_posts"]:
        body = post.get("text", "")
        project_terms = find_terms(body, PROJECT_KEYWORDS)
        codal_count = sum(1 for link in post.get("links", []) if is_codal_reference(urllib.parse.urljoin(post.get("source_page", TELEGRAM_BASE), link.get("url", ""))))
        attachment_count = sum(1 for link in post.get("links", []) if is_attachment_candidate(urllib.parse.urljoin(post.get("source_page", TELEGRAM_BASE), link.get("url", ""))))
        post_rows.append({
            "post_id": post.get("post_id", ""), "date": post.get("date", ""),
            "codal_count": codal_count, "attachment_count": attachment_count,
            "project_keywords": "؛ ".join(project_terms), "text": body,
            "permalink": post.get("permalink", ""),
        })
    write_sheet(
        wb, "پست‌های تلگرام",
        ["شناسه پست", "تاریخ", "تعداد لینک کدال", "تعداد پیوست اکسل", "کلیدواژه‌های طرح", "متن کامل پست", "لینک پست"],
        post_rows,
        ["post_id", "date", "codal_count", "attachment_count", "project_keywords", "text", "permalink"],
    )

    write_sheet(
        wb, "فهرست نمادها",
        ["نماد", "عنوان/متن لینک", "لینک فهرست گزارش‌های کدال", "منبع", "متن زمینه"],
        payload["symbol_lists"],
        ["symbol", "link_title", "url", "source", "source_context"],
    )
    write_sheet(
        wb, "پیوست‌های اکسل",
        ["لینک پیوست", "میزبان", "HTTP", "نوع محتوا", "حجم (بایت)", "XLSX معتبر؟", "تعداد شیت", "ردیف‌های خوانده‌شده", "نتیجه بررسی", "منبع"],
        payload["attachments"],
        ["url", "host", "http_status", "content_type", "body_bytes", "is_xlsx", "sheet_count", "extracted_rows", "parse_note", "source"],
    )
    write_sheet(
        wb, "داده‌های پیوست",
        ["لینک پیوست", "نام شیت", "شماره ردیف", "ردیف استخراج‌شده"],
        payload["attachment_rows"],
        ["attachment_url", "sheet_name", "row_number", "row_text"],
    )
    write_sheet(
        wb, "وضعیت منابع",
        ["نام منبع", "URL آزمایش‌شده", "HTTP", "نوع محتوا", "حجم پاسخ (بایت)", "آیتم‌های یافت‌شده", "یادداشت/خطا"],
        payload["source_status"],
        ["source_name", "url", "http_status", "content_type", "body_bytes", "items_found", "note"],
    )
    wb.save(OUT_XLSX)


def run(args):
    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    source_status = []
    all_candidates = []
    records = []
    symbol_lists = []

    listing = request_url(BOURSE_URL)
    listing_parser = PageAnchorParser()
    listing_page_text = ""
    if listing["body"]:
        listing_page_text = listing["body"].decode("utf-8", "replace")
        try:
            listing_parser.feed(listing_page_text)
        except Exception:
            pass
    bourse_codal_count = 0
    bourse_attachment_count = 0
    for link in listing_parser.links:
        url = urllib.parse.urljoin(BOURSE_URL, link["url"])
        if is_codal_reference(url):
            bourse_codal_count += 1
            title = link.get("title", "")
            context = filter_navigation_context(link.get("context") or "", title)
            row = make_record(url, title, context, "بورس‌تریدر", BOURSE_URL)
            if is_report_list(url):
                symbol_lists.append({
                    "symbol": row["symbol"], "link_title": row["link_title"],
                    "url": row["url"], "source": "بورس‌تریدر",
                    "source_context": row["source_context"],
                })
            else:
                records.append(row)
            all_candidates.append({"url": url, "source": "بورس‌تریدر", "context": context})
        if is_attachment_candidate(url):
            bourse_attachment_count += 1
            all_candidates.append({"url": url, "source": "بورس‌تریدر", "context": link.get("context", "")})
    source_status.append({
        "source_name": "بورس‌تریدر - صفحه اطلاعیه کدال", "url": BOURSE_URL,
        "http_status": listing["http_status"], "content_type": listing["content_type"],
        "body_bytes": listing["body_bytes"], "items_found": bourse_codal_count,
        "note": listing["error"] or ("response exceeded configured size limit" if listing["too_large"] else ""),
    })

    telegram_posts = fetch_telegram_pages(args.pages, source_status)
    telegram_detail_link_count = 0
    for post in telegram_posts:
        post_url = post.get("permalink") or TELEGRAM_BASE
        post_text = post.get("text", "")
        for link in post.get("links", []):
            raw_url = urllib.parse.urljoin(post.get("source_page", TELEGRAM_BASE), link.get("url", ""))
            if is_codal_reference(raw_url):
                if is_report_list(raw_url):
                    row = make_record(raw_url, link.get("title", ""), post_text, "تلگرام بورس‌تریدر", post_url, post.get("date", ""), post.get("post_id", ""))
                    symbol_lists.append({
                        "symbol": row["symbol"], "link_title": row["link_title"],
                        "url": row["url"], "source": "تلگرام بورس‌تریدر",
                        "source_context": post_text[:1800],
                    })
                else:
                    records.append(make_record(
                        raw_url, link.get("title", ""), post_text, "تلگرام بورس‌تریدر",
                        post_url, post.get("date", ""), post.get("post_id", ""),
                    ))
                    telegram_detail_link_count += 1
                all_candidates.append({"url": raw_url, "source": "تلگرام بورس‌تریدر", "context": post_text})
            if is_attachment_candidate(raw_url):
                all_candidates.append({"url": raw_url, "source": "تلگرام بورس‌تریدر", "context": post_text})

    announcements = merge_records(records)
    # Do not treat company-index links as announcement records.
    symbol_unique = {}
    for row in symbol_lists:
        key = row["url"]
        if key not in symbol_unique:
            symbol_unique[key] = row
        else:
            prior = symbol_unique[key]
            sources = set(prior.get("source", "").split("؛ ")) | set(row.get("source", "").split("؛ "))
            prior["source"] = "؛ ".join(sorted(s for s in sources if s))
            if len(row.get("source_context", "")) > len(prior.get("source_context", "")):
                prior["source_context"] = row["source_context"]
    symbol_lists = sorted(symbol_unique.values(), key=lambda item: item.get("symbol", ""))

    attachment_results, attachment_rows = download_xlsx_attachments(all_candidates, source_status)

    good_sources = sum(1 for row in source_status if row["http_status"] == 200 and row["body_bytes"] > 0)
    source_problem = any(row["http_status"] != 200 for row in source_status)
    if len(telegram_posts) == 0 and len(announcements) == 0 and len(symbol_lists) == 0:
        raise RuntimeError("No usable posts or Codal references parsed. Check source status and HTML selectors; blocked sources are not interpreted as zero activity.")

    quality_note = "لینک‌ها و متن‌های عمومی استخراج شدند؛ دسته‌بندی و ارقام خام خودکارند و برای تصمیم سرمایه‌گذاری باید با متن اطلاعیه/صورت مالی تطبیق داده شوند."
    if not announcements:
        quality_note += " هشدار: هیچ لینک گزارش جزئی استخراج نشد؛ فایل باید از نظر تغییر ساختار منبع بازبینی شود."
    if not attachment_rows:
        quality_note += " هیچ ردیف داده‌ای از پیوست XLSX قابل‌خواندن در این اجرا به دست نیامد؛ این به معنای نبود صورت مالی شرکت‌ها نیست."
    if source_problem:
        quality_note += " دست‌کم یک منبع/پیوست ناموفق بود؛ خطاهای آن در شیت «وضعیت منابع» آمده‌اند."
    summary = [
        {"metric": "زمان تولید UTC", "value": now_utc(), "note": "زمان ساخت فایل؛ نه تاریخ آخرین اطلاعیه"},
        {"metric": "صفحات تلگرام درخواست‌شده", "value": args.pages, "note": "حداکثر تعداد صفحه، هر صفحه معمولاً حدود ۲۰ پست"},
        {"metric": "پست‌های تلگرام استخراج‌شده", "value": len(telegram_posts), "note": "شناسه پست‌ها در صورت قابل‌تشخیص بودن حذف تکراری شده‌اند"},
        {"metric": "گزارش‌های کدال یکتا", "value": len(announcements), "note": "لینک‌های تکراری ادغام شده‌اند؛ گزارش‌ها لزوماً از صفحه مقصد خوانده نشده‌اند"},
        {"metric": "فهرست‌های نماد کدال", "value": len(symbol_lists), "note": "این لینک‌ها فهرست گزارش نماد هستند و با اطلاعیه منفرد یکی نیستند"},
        {"metric": "آیتم‌های طرح توسعه", "value": sum(1 for row in announcements if row.get("project_keywords")), "note": "فیلتر کلیدواژه‌ای؛ تأیید انسانی لازم است"},
        {"metric": "آیتم‌های مالی/ماهانه", "value": sum(1 for row in announcements if row.get("financial_keywords") or "صورت‌های مالی" in row.get("primary_category", "") or "ماهانه" in row.get("primary_category", "")), "note": "فیلتر کلیدواژه‌ای و نه تحلیل کامل صورت مالی"},
        {"metric": "پیوست‌های آزمایش‌شده", "value": len([a for a in attachment_results if a.get("http_status") != "SKIPPED"]), "note": f"حداکثر {MAX_ATTACHMENT_TESTS} پیوست در هر اجرا"},
        {"metric": "فایل‌های XLSX خوانده‌شده", "value": sum(1 for row in attachment_results if row.get("is_xlsx") == "بله"), "note": "محتوای خام در شیت «داده‌های پیوست» است"},
        {"metric": "منابع/درخواست‌های HTTP 200 با بدنه", "value": good_sources, "note": "شامل صفحات و درخواست‌های پیوست"},
        {"metric": "کیفیت/محدودیت‌ها", "value": "نیازمند بازبینی" if source_problem or not announcements or not attachment_rows else "استخراج اولیه موفق", "note": quality_note},
        {"metric": "API پولی", "value": "استفاده نشد", "note": "این اسکریپت از BRS_API_KEY استفاده نمی‌کند"},
    ]
    payload = {
        "generated_at_utc": now_utc(),
        "summary": summary,
        "source_status": source_status,
        "telegram_posts": telegram_posts,
        "announcements": announcements,
        "symbol_lists": symbol_lists,
        "attachments": attachment_results,
        "attachment_rows": attachment_rows,
        "quality_note": quality_note,
    }
    OUT_JSON.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    build_workbook(payload)
    print("EXTRACTION_SUMMARY")
    print(f"TELEGRAM_POSTS:{len(telegram_posts)}")
    print(f"CODAL_REPORT_LINKS_UNIQUE:{len(announcements)}")
    print(f"SYMBOL_REPORT_LISTS:{len(symbol_lists)}")
    print(f"PROJECT_KEYWORD_MATCHES:{sum(1 for row in announcements if row.get('project_keywords'))}")
    print(f"FINANCIAL_KEYWORD_MATCHES:{sum(1 for row in announcements if row.get('financial_keywords') or 'صورت‌های مالی' in row.get('primary_category','') or 'ماهانه' in row.get('primary_category',''))}")
    print(f"ATTACHMENTS_TESTED:{len([a for a in attachment_results if a.get('http_status') != 'SKIPPED'])}")
    print(f"XLSX_ATTACHMENTS_PARSED:{sum(1 for row in attachment_results if row.get('is_xlsx') == 'بله')}")
    print(f"OUTPUT_XLSX:{OUT_XLSX}")
    print(f"OUTPUT_JSON:{OUT_JSON}")
    print(f"QUALITY_NOTE:{quality_note}")
    return payload


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pages", type=int, default=DEFAULT_TELEGRAM_PAGES,
                        help="maximum Telegram preview pages to crawl (default: 8)")
    args = parser.parse_args()
    if args.pages < 1 or args.pages > 30:
        parser.error("--pages must be between 1 and 30")
    try:
        run(args)
    except Exception as exc:
        print(f"EXTRACTION_FAILED:{type(exc).__name__}:{exc}", file=sys.stderr)
        raise


if __name__ == "__main__":
    main()
